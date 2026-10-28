# /// script
# dependencies = ["openpyxl"]
# ///
"""Gera traducao/tecnicas/Tecnicas.xlsx: a planilha para traduzir os prefixos dos nomes de técnica.

Uso: uv run --script tools/ptbr/tecnicas_planilha.py [--forcar]

O jogo monta o nome como "<peça 1> <peça 2> <nome-base>" ("Sloth's Solid Água Espiritual", 仙磐灵水). As
peças vêm de traducao/tecnicas/BattleSkillPrefixName.json (cópia da tabela do kit): 200 da posição 1 e 200 da
posição 2. Em PT o nome vira "<nome-base> <adjetivo concordado> <complemento>" ("Água Espiritual Sólida do
Imortal"). Todas as combinações dariam milhões de linhas; a planilha tem uma linha por peça e uma aba de
combinações de teste com fórmulas, que mostram o resultado enquanto as peças são preenchidas.

Não sobrescreve uma planilha existente (já pode ter tradução) sem --forcar.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parents[2]
DIR = REPO / 'traducao/tecnicas'
OUT = DIR / 'Tecnicas.xlsx'
LT = REPO / 'traducao/LocalText'
LOG = Path('C:/Games/Tale of Immortal/ModExportData/Mod_6b2SYs/ptbr_tela_desconhecidas.txt')

AMARELO = PatternFill('solid', fgColor='FFF2CC')  # colunas para preencher
CINZA = PatternFill('solid', fgColor='EDEDED')  # colunas de fórmula
NEGRITO = Font(bold=True)

FEMININAS = {'Alma', 'Arte', 'Chave', 'Cicatriz', 'Corrente', 'Febre', 'Flor', 'Fonte', 'Fênix', 'Hipnose', 'Lápide',
             'Mão', 'Mãe', 'Mente', 'Morte', 'Neve', 'Noite', 'Parede', 'Pele', 'Ponte', 'Reunião', 'Sede', 'Serpente',
             'Torre', 'Fome'}
MASCULINAS = {'Coração', 'Dia', 'Fantasma', 'Guia', 'Mapa', 'Miasma', 'Plasma', 'Magma', 'Prisma', 'Sutra', 'Carma', 'Cometa',
              'Planeta', 'Poema', 'Tema', 'Enigma', 'Emblema', 'Diadema', 'Panda', 'Puma', 'Bagua'}
SINGULAR_EM_S = {'Lótus', 'Arco-íris', 'Cosmos', 'Quebra-cabeças', 'Atlas', 'Ônix'}


def genero_numero(nome):
    """Palpite pela primeira palavra; a aba Nomes-base deixa corrigir."""
    w = nome.split()[0]
    plural = w.endswith('s') and w not in SINGULAR_EM_S
    s = w
    if w.endswith('ões'):
        s = w[:-3] + 'ão'
    elif w.endswith(('res', 'zes')):
        s = w[:-2]
    elif plural:
        s = w[:-1]
    if s in MASCULINAS:
        fem = False
    elif s in FEMININAS:
        fem = True
    else:
        fem = s.endswith(('a', 'ade', 'ção', 'são', 'gem', 'ez', 'ude'))
    return ('f' if fem else 'm'), ('p' if plural else 's')


def ler(arq):
    return json.loads((LT / arq).read_text(encoding='utf-8'))


def nomes_base():
    """pt -> (ch, en, chave): os tipos genéricos, os passos e as técnicas de ataque."""
    bases = {}
    div = ler('_diversos.json')
    for k, v in sorted(div.items()):
        if re.fullmatch(r'(ability_name_last|stepName)\d+', k) and v['pt']:
            bases.setdefault(v['pt'], (v['ch'], v['en'], k))
    for k, v in sorted(ler('skill_name_middle.json').items()):
        pt = v['pt'].strip()
        if pt and pt != '???' and not re.search(r'[\u3400-\u9fff]', pt):
            en = v['en'] if not re.search(r'[\u3400-\u9fff]', v['en']) else ''
            bases.setdefault(pt, (v['ch'], en, k))
    return bases


def vistos_no_jogo(p1_en, p2_en):
    """Nomes de técnica do registro de tela do teste: (id1, id2, base, inglês como apareceu)."""
    if not LOG.exists():
        return []
    texto = LOG.read_text(encoding='utf-8')
    vistos = []
    for m in re.finditer(r"([A-Z][a-z]+'s?) ([A-Z][a-z-]+) ([^<\t\n]+?)(?=</color>|\t|\n)", texto):
        a, b, base = m.groups()
        if a in p1_en and b in p2_en:
            item = (p1_en[a], p2_en[b], base.strip(), m.group(0).strip())
            if item[:3] not in {v[:3] for v in vistos}:
                vistos.append(item)
    return vistos


def cabecalho(ws, titulos, larguras):
    for c, (t, w) in enumerate(zip(titulos, larguras), 1):
        cel = ws.cell(1, c, t)
        cel.font = NEGRITO
        cel.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    if OUT.exists() and '--forcar' not in sys.argv:
        sys.exit(f'{OUT} já existe (pode ter tradução). Use --forcar para gerar de novo.')
    pecas = json.loads((DIR / 'BattleSkillPrefixName.json').read_text(encoding='utf-8'))
    sentidos = json.loads((DIR / 'sentidos.json').read_text(encoding='utf-8'))
    obs = sentidos['obs']
    p1 = [p for p in pecas if p['index'] == '1']
    p2 = [p for p in pecas if p['index'] == '2']
    p1_en = {p['en']: int(p['id']) for p in p1}
    p2_en = {}
    for p in p2:
        p2_en.setdefault(p['en'], int(p['id']))
    por_id = {int(p['id']): p for p in pecas}

    wb = openpyxl.Workbook()
    leia = wb.active
    leia.title = 'Leia'
    for i, linha in enumerate([
        'Nomes de técnica gerados: o jogo junta peça 1 + peça 2 + nome-base.',
        'Chinês: 仙 + 磐 + 灵水 (imortal + rocha maciça + água espiritual). Inglês: "Sloth\'s Solid Spiritual Water".',
        'Em PT o mod vai montar: nome-base + adjetivo (peça 2, concordando) + complemento (peça 1).',
        'Exemplo: "Água Espiritual" + "Sólida" + "do Imortal" = "Água Espiritual Sólida do Imortal".',
        '',
        'O que preencher (colunas amarelas):',
        '  Peça 1: o complemento já com a preposição: "do Imortal", "da Lua", "do Dragão", "das Cinzas".',
        '  Peça 2: o adjetivo no masculino singular e, se mudar, no feminino ("Sólido" / "Sólida").',
        '          Feminino vazio = igual ao masculino ("Celeste", "Azul"). O plural o mod faz sozinho.',
        '          Se a peça 2 for uma locução ("de Névoa"), ela não varia: pode escrever no masculino só.',
        '  Nomes-base: gênero (m/f) e número (s/p) estão preenchidos por palpite; corrija os errados.',
        '',
        'Referência: "Sentido do hanzi" é o significado do caractere chinês. O inglês do jogo se perde da',
        'peça 1 nº 113 em diante (仙 "imortal" virou "Sloth\'s"); a coluna Obs. marca onde o inglês não bate.',
        '',
        'A aba Combinações mostra nomes completos com as fórmulas: as primeiras linhas são as técnicas vistas no',
        'teste do jogo; as outras passam por todas as 200 peças 1 e as 200 peças 2 pelo menos uma vez.',
        'Todas as combinações possíveis seriam 200 × 200 × ~700 nomes-base (milhões), por isso a planilha é por peça.',
    ], 1):
        leia.cell(i, 1, linha)
    leia.column_dimensions['A'].width = 120

    ws1 = wb.create_sheet('Peça 1')
    cabecalho(ws1, ['ID', 'Hanzi', 'Sentido do hanzi', 'Inglês do jogo', 'Obs.', 'Sua tradução (com "do/da/dos/das")',
                    'Exemplo'], [6, 7, 30, 16, 40, 34, 40])
    for r, p in enumerate(p1, 2):
        i = int(p['id'])
        nota = obs.get(str(i), obs['113-200'] if 113 <= i <= 200 else '')
        for c, v in enumerate([i, p['text'], sentidos[str(i)], p['en'], nota], 1):
            ws1.cell(r, c, v)
        ws1.cell(r, 6).fill = AMARELO
        ws1.cell(r, 7, f'=IF(F{r}="","","Técnica Sólida "&F{r})').fill = CINZA

    ws2 = wb.create_sheet('Peça 2')
    cabecalho(ws2, ['ID', 'Hanzi', 'Sentido do hanzi', 'Inglês do jogo', 'Obs.', 'Masculino', 'Feminino (se mudar)',
                    'Exemplo masculino', 'Exemplo feminino'], [6, 7, 30, 16, 34, 20, 20, 36, 36])
    for r, p in enumerate(p2, 2):
        i = int(p['id'])
        for c, v in enumerate([i, p['text'], sentidos[str(i)], p['en'], obs.get(str(i), '')], 1):
            ws2.cell(r, c, v)
        ws2.cell(r, 6).fill = AMARELO
        ws2.cell(r, 7).fill = AMARELO
        ws2.cell(r, 8, f'=IF(F{r}="","","Manual Secreto "&F{r}&" do Imortal")').fill = CINZA
        ws2.cell(r, 9, f'=IF(F{r}="","","Técnica "&IF(G{r}="",F{r},G{r})&" do Imortal")').fill = CINZA

    bases = nomes_base()
    ws3 = wb.create_sheet('Nomes-base')
    cabecalho(ws3, ['Nome-base (já traduzido)', 'Gênero (m/f)', 'Número (s/p)', 'Chinês', 'Inglês', 'Chave'],
              [40, 12, 12, 14, 30, 24])
    for r, (pt, (ch, en, chave)) in enumerate(bases.items(), 2):
        g, n = genero_numero(pt)
        for c, v in enumerate([pt, g, n, ch, en, chave], 1):
            ws3.cell(r, c, v)
        ws3.cell(r, 2).fill = AMARELO
        ws3.cell(r, 3).fill = AMARELO
    ultima = len(bases) + 1
    dv_g = DataValidation(type='list', formula1='"m,f"', allow_blank=False)
    dv_n = DataValidation(type='list', formula1='"s,p"', allow_blank=False)
    ws3.add_data_validation(dv_g)
    ws3.add_data_validation(dv_n)
    dv_g.add(f'B2:B{ultima}')
    dv_n.add(f'C2:C{ultima}')

    # Combinações: as vistas no jogo e uma passada por todas as peças.
    vistos = vistos_no_jogo(p1_en, p2_en)
    amostra = [b for b in dict.fromkeys(v[2] for v in vistos) if b in bases]
    amostra += [b for b in bases if bases[b][2].startswith(('ability_name_last', 'stepName')) and b not in amostra]
    linhas = [(a, b, base, ingles) for a, b, base, ingles in vistos if base in bases]
    for k in range(200):
        a, b = k + 1, 201 + (k * 73) % 200
        base = amostra[(k * 7) % len(amostra)]
        linhas.append((a, b, base, ''))
    ws4 = wb.create_sheet('Combinações')
    cabecalho(ws4, ['Como fica no jogo', 'Hanzi', 'Inglês do jogo', 'Peça 1', 'Peça 2', 'Nome-base',
                    'aux: peça 1', 'aux: adjetivo', 'aux: adjetivo concordado'], [52, 12, 44, 8, 8, 32, 22, 18, 22])
    for r, (a, b, base, ingles) in enumerate(linhas, 2):
        ch_base = bases[base][0]
        en_base = bases[base][1] or base
        ingles = ingles or f"{por_id[a]['en']} {por_id[b]['en']} {en_base}"
        ingles = re.sub(r'<[^>]+>', '', ingles)
        ws4.cell(r, 2, por_id[a]['text'] + por_id[b]['text'] + ch_base)
        ws4.cell(r, 3, ingles)
        ws4.cell(r, 4, a)
        ws4.cell(r, 5, b)
        ws4.cell(r, 6, base)
        ws4.cell(r, 7, f"=INDEX('Peça 1'!F:F,MATCH(D{r},'Peça 1'!A:A,0))&\"\"")
        genero = f"INDEX('Nomes-base'!B:B,MATCH(F{r},'Nomes-base'!A:A,0))"
        masc = f"INDEX('Peça 2'!F:F,MATCH(E{r},'Peça 2'!A:A,0))&\"\""
        fem = f"INDEX('Peça 2'!G:G,MATCH(E{r},'Peça 2'!A:A,0))&\"\""
        ws4.cell(r, 8, f'=IF(AND({genero}="f",{fem}<>""),{fem},{masc})')
        numero = f"INDEX('Nomes-base'!C:C,MATCH(F{r},'Nomes-base'!A:A,0))"
        h = f'H{r}'
        plural = (f'IF(RIGHT({h},1)="s",{h},IF(RIGHT({h},1)="l",LEFT({h},LEN({h})-1)&"is",'
                  f'IF(OR(RIGHT({h},1)="r",RIGHT({h},1)="z"),{h}&"es",IF(RIGHT({h},1)="m",LEFT({h},LEN({h})-1)&"ns",{h}&"s"))))')
        ws4.cell(r, 9, f'=IF(OR({h}="",ISNUMBER(SEARCH(" ",{h})),{numero}<>"p"),{h},{plural})')
        ws4.cell(r, 1, f'=IF(OR(G{r}="",I{r}=""),"(falta traduzir)",F{r}&" "&I{r}&" "&G{r})')
        ws4.cell(r, 1).fill = CINZA
        for c in (7, 8, 9):
            ws4.cell(r, c).font = Font(color='808080')

    DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'{OUT}: {len(p1)} peças 1, {len(p2)} peças 2, {len(bases)} nomes-base, '
          f'{len(linhas)} combinações ({len(vistos)} vistas no jogo)')


if __name__ == '__main__':
    main()
