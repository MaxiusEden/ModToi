# /// script
# dependencies = ["openpyxl"]
# ///
"""Planilha dos nomes de cidade -> ModAssets/ptbr/Cidades.json.

Uso: uv run --script tools/ptbr/cidades.py [planilha.xlsx]

A coluna "Como fica no jogo" ("Cidade do Herói Sem Nome") vira "do Herói Sem Nome": o mod põe na frente
o "Cidade" ou "Vila" que o jogo mostrou. Nome sem sufixo no jogo ("Cultivador Aliança") fica com o texto inteiro.
"""
import json
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / 'traducao/cidades/Cidades_Traduzidas_rodada5.xlsx'
OUT = REPO / 'ModProject_6b2SYs/ModProject/ModAssets/ptbr/Cidades.json'
COL_NOME, COL_JOGO = 6, 9  # "Nome no jogo" e "Como fica no jogo"
PREFIXO = 'Cidade '


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX
    ws = openpyxl.load_workbook(xlsx)['nomes_cidades']
    nomes, erros = {}, []
    for r in range(2, ws.max_row + 1):
        nome, jogo = ws.cell(r, COL_NOME).value, ws.cell(r, COL_JOGO).value
        if not nome:
            continue
        nome, jogo = str(nome).strip(), str(jogo or '').strip()
        if not jogo:
            erros.append(f'{nome}: sem "Como fica no jogo"')
            continue
        if ' ' in nome:
            nomes[nome] = jogo
        elif jogo.startswith(PREFIXO):
            nomes[nome] = jogo[len(PREFIXO):]
        else:
            erros.append(f'{nome}: "{jogo}" não começa com "{PREFIXO}"')
        if nome in nomes and any(not (c.isascii() or 'À' <= c <= 'ÿ') for c in nomes[nome]):
            erros.append(f'{nome}: caractere fora do ASCII e dos acentos latinos em "{nomes[nome]}"')
    repetidos = {v for v in nomes.values() if list(nomes.values()).count(v) > 1}
    erros += [f'tradução repetida: {v}' for v in sorted(repetidos)]
    if erros:
        print('\n'.join(erros))
        sys.exit(1)
    OUT.write_text(json.dumps(dict(sorted(nomes.items())), ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(f'{len(nomes)} nomes -> {OUT.relative_to(REPO)}')


if __name__ == '__main__':
    main()
