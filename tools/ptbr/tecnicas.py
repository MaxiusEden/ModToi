# /// script
# dependencies = ["openpyxl"]
# ///
"""Planilha dos nomes de técnica -> ModAssets/ptbr/Tecnicas.json.

Uso: uv run --script tools/ptbr/tecnicas.py [planilha.xlsx]

Saída: {"peca1": {id: "do Imortal"}, "peca2": {id: ["Sólido", "Sólida"]}, "bases": {"Água Espiritual": "fs"}}.
O mod troca o en da BattleSkillPrefixName pela peça 1 e pelo masculino da peça 2 e remonta o nome como
"<nome-base> <adjetivo concordado> <complemento>" (TechniqueNamePtBr.cs).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / 'traducao/tecnicas/Tecnicas.xlsx'
OUT = REPO / 'ModProject_6b2SYs/ModProject/ModAssets/ptbr/Tecnicas.json'
PERMITIDO = re.compile(r"^[ -~À-ÿ]+$")


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    wb = openpyxl.load_workbook(Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX)
    erros = []

    def texto(v, onde):
        v = (v or '').strip() if isinstance(v, str) or v is None else str(v)
        if v and not PERMITIDO.match(v):
            erros.append(f'{onde}: caractere fora do ASCII e dos acentos: {v!r}')
        return v

    peca1 = {}
    ws = wb['Peça 1']
    for r in range(2, ws.max_row + 1):
        i, pt = ws.cell(r, 1).value, texto(ws.cell(r, 6).value, f'Peça 1 linha {r}')
        if i is None:
            continue
        if not pt:
            erros.append(f'Peça 1 {i}: vazia')
        elif not re.match(r'(do|da|dos|das|de) ', pt):
            erros.append(f'Peça 1 {i}: não começa com do/da/dos/das/de: {pt!r}')
        peca1[str(i)] = pt

    peca2, fem_de = {}, {}
    ws = wb['Peça 2']
    for r in range(2, ws.max_row + 1):
        i = ws.cell(r, 1).value
        if i is None:
            continue
        m = texto(ws.cell(r, 6).value, f'Peça 2 linha {r}')
        f = texto(ws.cell(r, 7).value, f'Peça 2 linha {r}') or m
        if not m:
            erros.append(f'Peça 2 {i}: vazia')
        # O mod acha a peça pelo masculino: o mesmo masculino precisa ter o mesmo feminino.
        if fem_de.setdefault(m, f) != f:
            erros.append(f'Peça 2 {i}: "{m}" tem dois femininos ("{fem_de[m]}" e "{f}")')
        peca2[str(i)] = [m, f]

    bases = {}
    ws = wb['Nomes-base']
    for r in range(2, ws.max_row + 1):
        nome, g, n = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if not nome:
            continue
        if g not in ('m', 'f') or n not in ('s', 'p'):
            erros.append(f'Nomes-base linha {r} ({nome}): gênero/número inválido: {g!r}/{n!r}')
        bases[nome.strip()] = f'{g}{n}'

    if erros:
        print('\n'.join(erros))
        sys.exit(f'{len(erros)} erro(s); nada gravado.')
    OUT.write_text(json.dumps({'peca1': peca1, 'peca2': peca2, 'bases': bases}, ensure_ascii=False, indent=1) + '\n',
                   encoding='utf-8')
    print(f'{OUT}: {len(peca1)} peças 1, {len(peca2)} peças 2, {len(bases)} nomes-base')


if __name__ == '__main__':
    main()
